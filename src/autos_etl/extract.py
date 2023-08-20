from io import BytesIO

import openpyxl as xl
import pandas as pd
import requests
from docx import Document  # type: ignore

global EX1_FILE_A_URL
global EX1_FILE_B_URL
global EX1_FILE_C_URL
global EX2_FILE_A_URL

EX1_FILE_A_URL = "https://github.com/paulzuradzki/autos_db_modeling/raw/main/data/source/ex1_FileA.txt"
EX1_FILE_B_URL = "https://github.com/paulzuradzki/autos_db_modeling/raw/main/data/source/ex1__FileB.csv"
EX1_FILE_C_URL = "https://github.com/paulzuradzki/autos_db_modeling/raw/main/data/source/ex1__FileC.docx"
EX2_FILE_A_URL = "https://raw.githubusercontent.com/paulzuradzki/autos_db_modeling/main/data/source/ex2_FileA.xlsx"


def get_df_file_a() -> pd.DataFrame:
    def clean_price(value: str) -> str:
        return (
            value.replace("$", "").replace(",", "").replace(".00", "").replace('"', "")
        )

    response = requests.get(EX1_FILE_A_URL)
    file_obj = BytesIO(response.content)
    data_str = file_obj.read().decode("utf-8")
    data_lines = data_str.replace("\n\n", "\n").replace("\t\t", "\t").split("\n")

    data = [row.split("\t") for row in data_lines]

    column_names = [
        "id",
        "vin",
        "year",
        "make",
        "model",
        "trim",
        "wheel_drive_type",
        "color",
        "door_type",
        "drive_train_type",
        "msrp_amount",
    ]
    df = pd.DataFrame(data, columns=column_names)
    df.iloc[[6, 8, 9], 5:] = df.iloc[[6, 8, 9], 5:].shift(1, axis=1)

    df.loc[1, ["drive_train_type", "msrp_amount"]] = df.loc[1, "drive_train_type"].split(
        '"'
    )[  # type: ignore
        :-1
    ]

    df = (df
          .fillna("")
          
          # remove whitespace characters like \r and \r\r
          .applymap(lambda s: s.strip())
                    
          .assign(msrp_amount=lambda _df: _df["msrp_amount"]
                  .apply(clean_price)
                  )
          )

    return df


def get_df_file_b() -> pd.DataFrame:
    df_file_b = pd.read_csv(EX1_FILE_B_URL)
    return df_file_b


def get_df_file_c() -> pd.DataFrame:
    def _read_docx(file_path: str | BytesIO) -> str:
        document = Document(file_path)
        result = []

        for paragraph in document.paragraphs:
            result.append(paragraph.text)

        return "\n".join(result)

    response = requests.get(EX1_FILE_C_URL)
    file_obj = BytesIO(response.content)
    content: str = _read_docx(file_obj)
    dirty_data_items = [
        item.split("\n") for item in content.replace("\t", " ").split("\n\n")
    ]

    data_clean = []
    for item in dirty_data_items:
        name, address1, address2, *rest = item

        address = f"{address1} {address2}"
        notes = [x.strip() for x in rest]
        for note in notes:
            record = {"name": name, "address": address, "note": note}
            data_clean.append(record)

    df = pd.DataFrame(data_clean)
    return df


def get_df_ex2_file_a_transactions() -> pd.DataFrame:
    response = requests.get(EX2_FILE_A_URL)
    file_obj = BytesIO(response.content)

    wb = xl.load_workbook(file_obj)
    data = [[cell.value for cell in row] for row in wb["Sheet1"]["C3:P13"]]
    header = data[0]
    data = data[1:]
    data = [dict(zip(header, row)) for row in data]  # type: ignore
    df = pd.DataFrame(data).fillna("")
    return df


def get_df_ex2_file_a_customers() -> pd.DataFrame:
    response = requests.get(EX2_FILE_A_URL)
    file_obj = BytesIO(response.content)

    wb = xl.load_workbook(file_obj)
    data = [[cell.value for cell in row] for row in wb["Sheet1"]["C16:G26"]]
    header = data[0]
    data = data[1:]
    data = [dict(zip(header, row)) for row in data]  # type: ignore
    df = pd.DataFrame(data).fillna("")
    return df
